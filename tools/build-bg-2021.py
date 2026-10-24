#!/usr/bin/env python3
"""Bulgaria 2021 census -> the three GeoJSON files that pack.js turns into data/bg/2021/*.json

Inputs (keep them in raw/, which is not committed):
  bgr_admin2.geojson   265 municipality shapes (one per municipality, including Sarnitsa)
  bgr_admin1.geojson   28 oblast shapes
  bgr_admin0.geojson   country shape
  (the three boundary files share their border vertices exactly, so they can be simplified without gaps)
  bulgaria_population_detailed.xlsx               NSI InfoStat: population by oblast / municipality / settlement (English)
  bulgaria_population_detailed_in_bulgarian.xlsx  same table in Bulgarian
  bulgaria_ethnicity_detailed(.xlsx|_in_bulgarian.xlsx)   ethnic group by oblast / municipality
  bulgaria_religion_detailed(.xlsx|_in_bulgarian.xlsx)    religion by oblast / municipality
  minor_minorities.xlsx                           NSI: all ethnic groups, country level only
  (all six tables come from https://www.nsi.bg/infostat, census of 7 September 2021)

Usage (from the repo root; needs `pip install openpyxl numpy`):
  python3 tools/build-bg-2021.py raw build
  python3 tools/geo2topo.py build/bg_municipalities.geojson data/bg/2021/municipalities.json
  python3 tools/geo2topo.py build/bg_oblasts.geojson        data/bg/2021/oblasts.json
  python3 tools/geo2topo.py build/bg_country.geojson        data/bg/2021/country.json
(geo2topo.py does what `mapshaper -simplify 12% keep-shapes` + `node pack.js` do; those two still work on the build/*.geojson files too.)

Names. Every area carries two spellings: `n` (English = the official Bulgarian "Streamlined System" romanisation of the Bulgarian name,
Law on the Transliteration 2009, applied letter by letter, keeping the capital / small letters of the Cyrillic name: "Ново село" = "Novo selo")
and `nb` (Bulgarian, Cyrillic, without the "гр."/"с." prefix). The romanisation is computed here from the Cyrillic name of the NSI table; the English
column of the NSI table is only compared with it (differences are printed when the script runs). Nothing is taken over from the English names
of the shape files: they are only used to find the right shape for each municipality (they contain a few misspellings, see MATCH_FIX).

Only counts are stored (p, e, r). Percentages, majorities and shares are calculated in the browser.
In the NSI tables "-" means zero, so every "-" is stored as 0 (no withheld values -> "withheld": false in meta.json).
Categories (order = data/bg/2021/meta.json):
  ethnicity: Bulgarian, Turkish, Roma, the 15 smaller groups of minor_minorities.xlsx (largest first), Other, Cannot determine, Don't want to answer, Unknown
             The smaller groups were published for the whole country only: oblasts and municipalities have null for them (= not recorded) and
             their "Other" still contains those groups; only the country has them separately (its "Other" is the true remainder).
  religion:  Christian, Muslim, Jewish, Other, No religion, Cannot determine, Don't want to answer, Unknown
Settlements (EKATTE) only have a population figure (e = r = null); they are the breakdown `b` of each municipality.
"""
import json, sys, os, re
import openpyxl

raw, out = (sys.argv + ['raw', 'build'])[1:3]
os.makedirs(out, exist_ok=True)

# ---------------------------------------------------------------- romanisation (Streamlined System, 2009)
M = dict(zip('абвгдежзийклмнопрстуфхцчшщъьюя', ['a','b','v','g','d','e','zh','z','i','y','k','l','m','n','o','p','r','s','t','u','f','h','ts','ch','sh','sht','a','y','yu','ya']))
def tr(s):
    out_ = []
    for w in re.split(r'(\s+|-)', s):
        if not w or re.fullmatch(r'\s+|-', w): out_.append(w); continue
        x = re.sub('ия$', '\0', w.lower()).replace('ьо', 'yo')          # word-final -ия is -ia, "ьо" is "yo"
        r = ''.join('ia' if c == '\0' else M.get(c, c) for c in x)
        out_.append(r[:1].upper() + r[1:] if w[:1].isupper() else r)    # a word starts with a capital letter only if the Cyrillic word does
    return ''.join(out_)

# NSI typos in the English table, corrected by the romanisation of the Bulgarian name: 47319 "Markovi" -> Марково = Markovo; 22753 "Dolno  Osenovo" (double space)
# Two oblasts are both called "София" in Bulgarian: the oblast (Sofia) and "София (столица)" = Sofia (stolitsa), the capital's oblast. Municipality "Столична" = Stolichna.
# 72655 Томпсън = Tompsan (the village is named after Thompson, but the Bulgarian name is romanised like any other).
NAME_FIX = {}     # code -> English name, only for names that must NOT be the plain romanisation (none)

# Shape files -> census: the shape file's municipality (by pcode) whose name does not resemble the census name closely enough
MATCH_FIX = {'BG003003': 'DOB15', 'BG021001': 'SOF46'}          # "Dobrichka" is Добрич-селска (Dobrich-selska); "Dobrich" (BG003002) is the town municipality
OBL_FIX = {'Sofia': 'SFO', 'Sofia-City': 'SOF'}   # shape file oblast name -> census code where the names differ

# ---------------------------------------------------------------- read the tables
def lvl(c):
    c = str(c)
    return ('settle' if re.fullmatch(r'\d{5}', c) else 'muni' if re.fullmatch(r'[A-Z]{3}\d{2}', c) else 'obl' if re.fullmatch(r'[A-Z]{3}', c)
            else 'reg' if re.fullmatch(r'BG\d{2}', c) else 'mac' if re.fullmatch(r'BG\d', c) else 'bg' if c == 'BG' else None)
def num(v):
    if v == '-': return 0
    if isinstance(v, (int, float)): return int(v)
    raise ValueError(f'unexpected cell {v!r}')
def read(f):
    ws = openpyxl.load_workbook(os.path.join(raw, f + '.xlsx'), read_only=True).active
    return [r for r in ws.iter_rows(values_only=True) if r[0] is not None and lvl(r[0])]
pop_en, pop_bg = read('bulgaria_population_detailed'), read('bulgaria_population_detailed_in_bulgarian')
eth_en, eth_bg = read('bulgaria_ethnicity_detailed'), read('bulgaria_ethnicity_detailed_in_bulgarian')
rel_en, rel_bg = read('bulgaria_religion_detailed'), read('bulgaria_religion_detailed_in_bulgarian')
assert [r[0] for r in pop_en] == [r[0] for r in pop_bg] and all(a[2] == b[2] for a, b in zip(pop_en, pop_bg))
for a, b in ((eth_en, eth_bg), (rel_en, rel_bg)):
    assert [r[0] for r in a] == [r[0] for r in b] and [r[2:] for r in a] == [r[2:] for r in b], 'English and Bulgarian tables differ'
NE, NR = 7, 8
eth = {r[0]: (num(r[2]), [num(x) for x in r[3:3 + NE]]) for r in eth_en if lvl(r[0]) in ('bg', 'obl', 'muni')}
rel = {r[0]: (num(r[2]), [num(x) for x in r[3:3 + NR]]) for r in rel_en if lvl(r[0]) in ('bg', 'obl', 'muni')}
for k, (p, c) in eth.items(): assert sum(c) == p, f'ethnicity {k} does not add up'
for k, (p, c) in rel.items(): assert sum(c) == p, f'religion {k} does not add up'

PRE = re.compile(r'^(гр\.|с\.|ман\.)\s*')
oblasts, munis, order = {}, {}, []
cur_o = cur_m = None
problems = []
for a, b in zip(pop_en, pop_bg):
    code, L, p = a[0], lvl(a[0]), num(a[2])
    if L == 'obl':
        cur_o = oblasts[code] = dict(code=code, p=p, bg=b[1], munis=[])
    elif L == 'muni':
        cur_m = munis[code] = dict(code=code, o=cur_o['code'], p=p, bg=b[1], nsi=a[1], settle=[])
        cur_o['munis'].append(code)
    elif L == 'settle':
        bgn = PRE.sub('', b[1]); name = NAME_FIX.get(code) or tr(bgn)
        nsi = re.sub(r'\s+', ' ', a[1])
        if name.lower() != nsi.lower(): problems.append((code, nsi, bgn, name))
        cur_m['settle'].append(dict(n=name, nb=bgn, p=p, e=None, r=None, **({'y': 't'} if b[1].startswith('гр.') else {})))
for code, m in munis.items():
    m['n'] = NAME_FIX.get(code) or tr(m['bg'])
    if m['n'].lower() != m['nsi'].lower(): problems.append((code, m['nsi'], m['bg'], m['n']))
    assert m['p'] == eth[code][0] == rel[code][0] == sum(s['p'] for s in m['settle']), f'{code}: population differs between tables'
for code, o in oblasts.items():
    o['n'] = NAME_FIX.get(code) or tr(o['bg'])
    assert o['p'] == eth[code][0] == rel[code][0] == sum(munis[m]['p'] for m in o['munis'])
    for t in (eth, rel):
        assert all(t[code][1][i] == sum(t[m][1][i] for m in o['munis']) for i in range(len(t[code][1]))), code
tot = pop_en[0]; assert tot[0] == 'BG' and tot[2] == eth['BG'][0] == rel['BG'][0] == sum(o['p'] for o in oblasts.values())
print(f'{len(oblasts)} oblasts, {len(munis)} municipalities, {sum(len(m["settle"]) for m in munis.values())} settlements, population {tot[2]:,}')
print('romanisation: names where our spelling differs from the NSI English table (apart from capitals):')
for x in problems: print('   ', x)

# smaller ethnic groups, country level only (minor_minorities.xlsx); English names as in the file, "Ucrainian" corrected
mm = [r for r in openpyxl.load_workbook(os.path.join(raw, 'minor_minorities.xlsx'), read_only=True).active.iter_rows(values_only=True) if r[1] and isinstance(r[2], int)]
mm = {r[1].replace('Ucrainian', 'Ukrainian'): r[2] for r in mm}
assert mm.pop('Total') == eth['BG'][0] and [mm.pop(k) for k in ('Bulgarian', 'Turkish', 'Roma')] == eth['BG'][1][:3]
assert [mm.pop(k) for k in ('Cannot determine', "Don't want to answer", 'Unknown2')] == eth['BG'][1][4:]
MINOR_OTHER = mm.pop('Other')
MINORS = sorted(mm, key=lambda k: -mm[k])                 # 15 groups, largest first = their order in meta.json
assert len(MINORS) == 15 and sum(mm.values()) + MINOR_OTHER == eth['BG'][1][3], 'the smaller groups plus Other must equal the NSI Other of the country'
def eth7(code, e):    # 7 categories of the oblast / municipality tables -> the full category list
    mins = [mm[k] for k in MINORS] if code == 'BG' else [None] * len(MINORS)
    return e[:3] + mins + [MINOR_OTHER if code == 'BG' else e[3]] + e[4:]
def area(code):
    d = dict(p=eth[code][0], e=eth7(code, eth[code][1]), r=rel[code][1])
    return d

# ---------------------------------------------------------------- shapes
import difflib
def geo(f): return json.load(open(os.path.join(raw, f), encoding='utf-8'))['features']
OB_BG = {'Софийска': 'София', 'София-град': 'София (столица)'}           # the shape file's Bulgarian oblast names where they differ from the census
obl_by_bg = {}
for c, o in oblasts.items(): obl_by_bg.setdefault(o['bg'], []).append(c)
shape_of = {}
for f in geo('bgr_admin2.geojson'):
    p = f['properties']; pc = p['adm2_pcode']
    if pc in MATCH_FIX: code = MATCH_FIX[pc]
    else:
        ob = obl_by_bg[OB_BG.get(p['adm1_bg'], p['adm1_bg'])]
        cand = sorted(((difflib.SequenceMatcher(None, p['adm2_name'].lower(), munis[c]['n'].lower()).ratio(), c) for c in munis if munis[c]['o'] in ob), reverse=True)
        assert cand[0][0] > 0.75 and cand[0][0] - (cand[1][0] if len(cand) > 1 else 0) > 0.1, f'unclear match for {p["adm2_name"]}: {cand[:2]}'
        code = cand[0][1]
    assert code not in shape_of, f'{code} matched twice'
    shape_of[code] = f['geometry']
assert set(shape_of) == set(munis), set(munis) - set(shape_of)            # every municipality (Sarnitsa included) has a shape of its own

feats = []
for code, m in munis.items():
    o = oblasts[m['o']]
    feats.append(dict(type='Feature', geometry=shape_of[code], properties=dict(n=m['n'], nb=m['bg'], j=o['n'], jb=o['bg'], **area(code), b=m['settle'])))
assert len(feats) == 265
json.dump(dict(type='FeatureCollection', features=feats), open(os.path.join(out, 'bg_municipalities.geojson'), 'w'), ensure_ascii=False)

feats = []
for f in geo('bgr_admin1.geojson'):
    nm = f['properties']['adm1_name']
    code = OBL_FIX.get(nm) or [c for c, o in oblasts.items() if o['n'] == nm][0]
    o = oblasts[code]
    feats.append(dict(type='Feature', geometry=f['geometry'], properties=dict(n=o['n'], nb=o['bg'], **area(code), b=[dict(n=munis[c]['n'], nb=munis[c]['bg'], **area(c)) for c in o['munis']])))
assert len(feats) == 28 and len({f['properties']['n'] for f in feats}) == 28
json.dump(dict(type='FeatureCollection', features=feats), open(os.path.join(out, 'bg_oblasts.geojson'), 'w'), ensure_ascii=False)

c = geo('bgr_admin0.geojson')[0]
b = [dict(n=o['n'], nb=o['bg'], **area(k)) for k, o in oblasts.items()]
json.dump(dict(type='FeatureCollection', features=[dict(type='Feature', geometry=c['geometry'],
    properties=dict(n='Bulgaria', nb='България', **area('BG'), b=b))]), open(os.path.join(out, 'bg_country.geojson'), 'w'), ensure_ascii=False)
print('wrote', out)
