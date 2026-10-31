#!/usr/bin/env python3
"""North Macedonia 2021 census -> the two GeoJSON files that pack.js turns into data/mk/2021/*.json

Inputs (keep them in raw/, which is not committed):
  north_mkd_ethnic.xlsx      MakStat (makstat.stat.gov.mk/PXWeb/pxweb/mk/MakStat/) table T1008P21: resident population by ethnic affiliation, by municipality
  north_mkd_religion.xlsx    MakStat table T1012P21: resident population by religious affiliation, by municipality
  north_mkd_population.xlsx  MakStat table T1019P21: resident population by municipality (only used to double-check the totals)
  geoBoundaries-MKD-ADM2_obshtina_municipalities.geojson   84 municipality shapes (the pre-2013 map, see below)

Usage (from the repo root; needs `pip install openpyxl shapely numpy`, and `cd tools && npm install`):
  python3 tools/build-mk-2021.py raw build
  cd tools
  node pack.js ../build/mk_municipalities.geojson ../data/mk/2021/municipalities.json
  node pack.js ../build/mk_country.geojson        ../data/mk/2021/country.json

Municipalities. The census tables have the 80 municipalities of the 2013 reform (the 10 municipalities of the City of Skopje included; "Град Скопје"
is only their total). The geoBoundaries file still has 84: Drugovo, Oslomej, Zajas and Vraneštica were merged into Kičevo in March 2013, so their four shapes
are joined to Kičevo's (Kičevo: 56,734 residents in 2002 within the new borders, 39,669 in 2021). The country shape is the union of the 80 municipalities.

Names. The tables are in Macedonian (Cyrillic); n is the usual English spelling, made from the Cyrillic name letter by letter:
ж ч ш = ž č š, ѓ = gj, ќ = kj, џ = dž, ѕ = dz, љ = lj, њ = nj, ј = j, ц = c, х = h ("Чучер - Сандево" = Čučer-Sandevo); only "Маврово и Ростуше" is written
"Mavrovo and Rostuša" (English name). The shapes are found by comparing the names romanised the geoBoundaries way (ц = ts, ж = zh ...); three need help (OVERRIDE).
The ten Skopje municipalities have j = "City of Skopje"; in the country file they are nested (c) under one "City of Skopje" row.

Categories (order = data/mk/2021/meta.json)
  ethnicity: Macedonian, Albanian, Turkish, Roma, Vlach, Serbian, Bosniak, Other, then Not declared, Unknown, Data from administrative sources
  religion:  Orthodox, Muslim, Catholic, Christian (denomination not specified), Protestant, Evangelical, Evangelical-Methodist, Jehovah’s Witnesses,
             Atheist (not a believer), Other, then Not declared, Unknown, Data from administrative sources
"Data from administrative sources" = "Лица за кои податоците се превземени од административни извори": 132,260 people (7.2%) for whom the census had no
interview and so no ethnicity or religion. In these tables "-" means zero, so every "-" is stored as 0 (no withheld values -> "withheld": false in meta.json).
Only counts are stored (p, e, r); percentages, majorities and shares are calculated in the browser.
"""
import json, sys, os, difflib
import numpy as np
import openpyxl
from shapely.geometry import shape, Polygon, MultiPolygon
from shapely.geometry.polygon import orient
from shapely.ops import unary_union
from shapely import make_valid

raw, out = (sys.argv + ['raw', 'build'])[1:3]
os.makedirs(out, exist_ok=True)

# ---------------------------------------------------------------- tables
def num(v):
    if v == '-': return 0
    if isinstance(v, (int, float)): return int(v)
    raise ValueError(f'unexpected cell {v!r}')
def table(f, ncat):
    rr = list(openpyxl.load_workbook(os.path.join(raw, f), read_only=True).active.iter_rows(values_only=True))
    res = []
    for r in rr[4:]:
        if r[0] is None: break                                    # footnotes start after an empty row
        res.append(dict(name=str(r[0]).strip(), p=num(r[1]), c=[num(x) for x in r[2:2 + ncat]]))
    return res
eth, rel = table('north_mkd_ethnic.xlsx', 11), table('north_mkd_religion.xlsx', 13)
pop = {r[0].strip(): num(r[2]) for r in openpyxl.load_workbook(os.path.join(raw, 'north_mkd_population.xlsx'), read_only=True).active.iter_rows(min_row=5, values_only=True) if r[0] and r[1] == 'Возрасна група - ВКУПНО'}
assert [(a['name'], a['p']) for a in eth] == [(a['name'], a['p']) for a in rel] and len(eth) == 82, 'ethnicity and religion tables do not line up'
for a in eth + rel: assert sum(a['c']) == a['p'], f"{a['name']}: categories do not add up"
assert all(pop[a['name']] == a['p'] for a in eth), 'population differs from T1019P21'
country, skopje = eth[0], eth[1]; cr = rel[0]
assert country['name'] == 'Република Северна Македонија' and skopje['name'] == 'Град Скопје'
muni_e, muni_r = eth[2:], rel[2:]
assert len(muni_e) == 80 and sum(a['p'] for a in muni_e) == country['p'] == 1836713
SK = muni_e[:10]; assert sum(a['p'] for a in SK) == skopje['p'] and [a['name'] for a in SK][0] == 'Аеродром'

# ---------------------------------------------------------------- names
LAT = dict(zip('абвгдѓежзѕијклљмнњопрстќуфхцчџш', ['a','b','v','g','d','gj','e','ž','z','dz','i','j','k','l','lj','m','n','nj','o','p','r','s','t','kj','u','f','h','c','č','dž','š']))
def latin(s):
    if s == 'Маврово и Ростуше': return 'Mavrovo and Rostuša'
    out_ = []
    for w in s.replace(' - ', '-').split(' '):
        parts = [''.join(LAT[c] for c in x.lower()) for x in w.split('-')]
        out_.append('-'.join(p[:1].upper() + p[1:] for p in parts))
    return ' '.join(out_)
GB = dict(zip('абвгдѓежзѕијклљмнњопрстќуфхцчџш', ['a','b','v','g','d','gj','e','zh','z','dz','i','j','k','l','lj','m','n','nj','o','p','r','s','t','kj','u','f','h','ts','ch','dzh','sh']))
gb = lambda s: ''.join(GB.get(c, c) for c in s.lower().replace(' - ', ' ').replace('-', ' '))
OVERRIDE = {'Центар': 'Centar', 'Центар Жупа': 'Centar Zhupa', 'Маврово и Ростуше': 'Mavrovo and Rostusha'}
MERGED_INTO_KICEVO = ['Drugovo', 'Oslomej', 'Zajas', 'Vraneshtitsa']

# ---------------------------------------------------------------- shapes
feats = json.load(open(os.path.join(raw, 'geoBoundaries-MKD-ADM2_obshtina_municipalities.geojson'), encoding='utf-8'))['features']
assert len(feats) == 84
geo = {}
for f in feats:
    g = shape(f['geometry']); geo[f['properties']['shapeName']] = g if g.is_valid else make_valid(g)
shape_of = {}
for a in muni_e:
    if a['name'] in OVERRIDE: s = OVERRIDE[a['name']]
    else:
        c = sorted(((difflib.SequenceMatcher(None, gb(a['name']), x.lower().replace(' - ', ' ').replace('-', ' ')).ratio(), x) for x in geo), reverse=True)
        assert c[0][0] > 0.9 and c[0][0] - c[1][0] > 0.1, f"unclear shape for {a['name']}: {c[:2]}"
        s = c[0][1]
    assert s not in shape_of.values(), f'{s} matched twice'
    shape_of[a['name']] = s
assert set(geo) - set(shape_of.values()) == set(MERGED_INTO_KICEVO), set(geo) - set(shape_of.values())

polys = lambda g: list(g.geoms) if g.geom_type in ('MultiPolygon', 'GeometryCollection') else [g]
def only_polygons(g):
    ps = [p for p in polys(g) if p.geom_type == 'Polygon' and not p.is_empty]
    return ps[0] if len(ps) == 1 else MultiPolygon(ps)
def dissolve(geoms, fill=1e-5):
    """union of the shapes under an area; holes smaller than `fill` deg2 (slivers between shapes that do not quite touch) are filled"""
    u = unary_union(geoms)
    return only_polygons(unary_union([Polygon(p.exterior, [r for r in p.interiors if Polygon(r).area >= fill]) for p in polys(u) if p.geom_type == 'Polygon']))
def ring_xy(r, nd=6):
    a = np.round(np.asarray(r.coords), nd)
    return a[np.concatenate([[True], (a[1:] != a[:-1]).any(axis=1)])]
def geometry(g):
    res = []
    for p in polys(g):
        p = orient(p, 1.0)
        rings = [r.tolist() for r in [ring_xy(p.exterior)] + [ring_xy(r) for r in p.interiors] if len(r) >= 4]
        if rings and len(rings[0]) >= 4: res.append(rings)
    assert res
    return {'type': 'MultiPolygon', 'coordinates': res} if len(res) > 1 else {'type': 'Polygon', 'coordinates': res[0]}
def write(name, fs):
    path = os.path.join(out, name + '.geojson')
    json.dump({'type': 'FeatureCollection', 'features': [{'type': 'Feature', 'properties': p, 'geometry': geometry(g)} for p, g in fs]}, open(path, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print('wrote', path, len(fs), 'features')

# ---------------------------------------------------------------- municipalities and country
sk_names = {a['name'] for a in SK}
fs, rows_ = [], []
for a, b in zip(muni_e, muni_r):
    g = geo[shape_of[a['name']]]
    if shape_of[a['name']] == 'Kichevo': g = dissolve([g] + [geo[x] for x in MERGED_INTO_KICEVO], fill=1e-6)
    p = dict(n=latin(a['name']), p=a['p'], e=a['c'], r=b['c'])
    if a['name'] in sk_names: p['j'] = 'City of Skopje'
    fs.append((p, g)); rows_.append(p)
write('mk_municipalities', fs)
cb = [dict(n=p['n'], p=p['p'], e=p['e'], r=p['r']) for p in rows_ if 'j' not in p]
sk_row = dict(n='City of Skopje', p=skopje['p'], e=eth[1]['c'], r=rel[1]['c'], c=[dict(n=p['n'], p=p['p'], e=p['e'], r=p['r']) for p in rows_ if 'j' in p])
assert len(cb) == 70 and sum(u['p'] for u in sk_row['c']) == sk_row['p']
write('mk_country', [(dict(n='North Macedonia', p=country['p'], e=country['c'], r=cr['c'], b=[sk_row] + cb), dissolve([g for _, g in fs]))])
print(f"{len(fs)} municipalities, population {country['p']:,}")
